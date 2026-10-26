from pandas.errors import ClosedFileError
import streamlit as st
import pandas as pd
import plotly.express as px
import os
import json
import base64
from datetime import datetime
from io import BytesIO
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# --- Configuration & Styling ---
FEEDBACK_DIR = "feedback"
CONFIG_DIR = "config"
os.makedirs(FEEDBACK_DIR, exist_ok=True)
os.makedirs(CONFIG_DIR, exist_ok=True)

CATEGORIES_JSON = os.path.join(CONFIG_DIR, "categories.json")

# Default Category Mapping Rules
DEFAULT_CATEGORY_MAPPING = {
    'Boxer': ['boxer'],
    'Tank Top': ['tank top', 'tanktop', 'tank', 'top'],
    'Cuban Shirt': ['Cuban Collar Shirt', 'cuban'],
    'Panjabi': ['panjabi', 'punjabi'],
    'Jeans': ['jeans'],
    'Formal Shirt': ['executive', 'formal'],
    'Denim Shirt': ['denim shirt'],
    'Flannel Shirt': ['flannel'],
    'Polo Shirt': ['polo'],
    'Cargo Trousers': ['cargo', 'cargo trousers'],
    'Trousers': ['trousers', 'pant', 'trcargo', 'trouser', 'joggers', 'track pant', 'jogger'],
    'Twill': ['twill chino', 'twill', 'chino'],
    'Mask': ['mask'],
    'Water Bottle': ['water bottle'],
    'Contrast Shirt': ['contrast'],
    'Turtleneck': ['turtleneck', 'mock neck'],
    'Drop Shoulder': ['drop', 'shoulder'],
    'Wallet': ['wallet'],
    'Card Holder': ['card holder'],
    'Passport Holder': ['passport holder'],
    'Kaftan Shirt': ['kaftan'],
    'Active Wear': ['active wear'],
    'Jersy': ['jersy'],
    'Sweatshirt': ['sweatshirt', 'hoodie', 'pullover'],
    'Jacket': ['jacket', 'outerwear', 'coat'],
    'Belt': ['belt'],
    'Sweater': ['sweater', 'cardigan', 'knitwear'],
    'Cap': ['cap'],
    'Leather Bag': ['bag', 'backpack'],
    'Mug': ['mug'],
    'Hoodie': ['hoodie'],
    'Knit Cap': ['knit cap'],
}

def load_category_mapping():
    """Loads custom categories from config/categories.json if available, otherwise defaults."""
    if os.path.exists(CATEGORIES_JSON):
        try:
            with open(CATEGORIES_JSON, "r", encoding="utf-8") as f:
                data = json.load(f)
                if isinstance(data, dict) and data:
                    return data
        except Exception:
            pass
    return DEFAULT_CATEGORY_MAPPING

CATEGORY_MAPPING = load_category_mapping()

# Bright, vibrant qualitative color palette for high-contrast synchronized chart visuals
CATEGORY_PALETTE = list(dict.fromkeys(
    px.colors.qualitative.Plotly +
    px.colors.qualitative.Vivid +
    px.colors.qualitative.Bold +
    px.colors.qualitative.D3 +
    px.colors.qualitative.Set1 +
    px.colors.qualitative.Dark24
))

ALL_KNOWN_CATEGORIES = list(CATEGORY_MAPPING.keys()) + [
    'FS T-Shirt', 'T-Shirt', 'FS Shirt', 'HS Shirt', 'Others', 'Unknown'
]

CATEGORY_COLOR_MAP = {
    cat: CATEGORY_PALETTE[i % len(CATEGORY_PALETTE)] for i, cat in enumerate(ALL_KNOWN_CATEGORIES)
}

def get_category_color_map(categories):
    """Returns a synchronized category-to-color mapping ensuring identical colors across charts."""
    color_map = {}
    known_cats = list(CATEGORY_COLOR_MAP.keys())
    for i, cat in enumerate(categories):
        if cat in CATEGORY_COLOR_MAP:
            color_map[cat] = CATEGORY_COLOR_MAP[cat]
        else:
            color_map[cat] = CATEGORY_PALETTE[(len(known_cats) + i) % len(CATEGORY_PALETTE)]
    return color_map

LOGO_PNG = "assets/deen_logo.png"

def load_logo():
    """Loads logo as base64 string from assets/deen_logo.png."""
    if os.path.exists(LOGO_PNG):
        try:
            with open(LOGO_PNG, "rb") as f:
                return base64.b64encode(f.read()).decode()
        except: pass
    return ""

def apply_custom_styles():
    """Applies premium CSS styles to the dashboard."""
    st.markdown("""
        <style>
        /* Base Container */
        .main { background-color: transparent; }
        
        /* Metric Styling - Premium Look */
        div[data-testid="stMetric"] { 
            background-color: var(--background-secondary, #ffffff); 
            padding: 24px; 
            border-radius: 16px; 
            box-shadow: 0 4px 20px rgba(0,0,0,0.06);
            border: 1px solid var(--secondary-background-color, #ececec);
            transition: all 0.3s cubic-bezier(0.4, 0, 0.2, 1);
        }
        
        div[data-testid="stMetric"]:hover {
            transform: translateY(-4px);
            box-shadow: 0 12px 24px rgba(0,0,0,0.12);
            border-color: #007bff;
        }

        /* Metric Label Support */
        [data-testid="stMetricLabel"] {
            font-weight: 600 !important;
            font-size: 1.1rem !important;
            margin-bottom: 8px !important;
            color: var(--text-color, #1a1a1b) !important;
        }

        /* Metric Value - Ensure Full Visibility */
        [data-testid="stMetricValue"], 
        [data-testid="stMetricValue"] > div {
            font-weight: 800 !important;
            font-size: 1.5rem !important; /* Reduced for better fit on large numbers */
            color: var(--text-color, #1a1a1b) !important;
            white-space: nowrap !important;
            overflow: visible !important;
            text-overflow: clip !important;
            min-width: fit-content !important;
        }

        div[data-testid="stMetric"] {
            overflow: visible !important;
        }

        /* Dark Mode Text Consistency - Users specifically want white text */
        @media (prefers-color-scheme: dark) {
            [data-testid="stMetricLabel"], 
            [data-testid="stMetricValue"],
            [data-testid="stMetricValue"] > div { 
                color: #ffffff !important; 
            }
            div[data-testid="stMetric"] {
                background-color: #1e1e1e !important;
                border-color: #333 !important;
            }
        }

        /* Streamlit Dark Theme Detection */
        [data-theme="dark"] [data-testid="stMetricLabel"],
        [data-theme="dark"] [data-testid="stMetricValue"],
        [data-theme="dark"] [data-testid="stMetricValue"] > div {
            color: #ffffff !important;
        }
        
        [data-theme="dark"] div[data-testid="stMetric"] {
            background-color: #1e1e1e !important;
            box-shadow: 0 4px 20px rgba(0,0,0,0.4);
            border: 1px solid #333;
        }

        .stButton>button { 
            width: 100%; 
            border-radius: 12px; 
            height: 3.5em; 
            background: linear-gradient(135deg, #007bff 0%, #0056b3 100%); 
            color: white; 
            font-weight: 600;
            border: none;
            transition: all 0.3s ease;
        }
        .stButton>button:hover {
            transform: scale(1.02);
            background: linear-gradient(135deg, #0056b3 0%, #004085 100%);
            box-shadow: 0 6px 12px rgba(0,0,0,0.15);
        }
        div[data-testid="stExpander"] { 
            border: 1px solid var(--secondary-background-color, #eee); 
            box-shadow: 0 4px 12px rgba(0,0,0,0.05);
            background: var(--background-secondary, white);
            border-radius: 14px;
            margin-bottom: 1.5rem;
            padding: 5px;
            transition: border-color 0.3s ease;
        }
        
        [data-theme="dark"] div[data-testid="stExpander"] {
            background-color: #1a1a1a !important;
            border-color: #333 !important;
            box-shadow: 0 4px 15px rgba(0,0,0,0.3);
        }

        @media (prefers-color-scheme: dark) {
            div[data-testid="stExpander"] {
                background-color: #1a1a1a !important;
                border-color: #333 !important;
            }
        }
        .sticky-footer {
            position: fixed;
            left: 0;
            bottom: 0;
            width: 100%;
            background-color: #ffffff; /* Pure white in light mode */
            color: #000000 !important; /* Pure black in light mode */
            text-align: center;
            padding: 18px 0;
            border-top: 1px solid #e9ecef;
            z-index: 999;
            font-size: 0.9rem;
            backdrop-filter: blur(12px);
            box-shadow: 0 -4px 30px rgba(0,0,0,0.06);
        }
        [data-theme="dark"] .sticky-footer {
            background-color: #121212 !important; /* Pure dark in dark mode */
            color: #ffffff !important;
            border-top: 1px solid #333 !important;
        }
        @media (prefers-color-scheme: dark) {
            .sticky-footer {
                background-color: #121212 !important;
                color: #ffffff !important;
            }
        }
        .footer-content-inner {
            display: flex;
            align-items: center;
            justify-content: center;
            gap: 6px;
        }
        .brand-wrapper {
            display: flex;
            align-items: center;
            gap: 2px;
        }
        .small-logo {
            height: 24px;
            width: auto;
            filter: grayscale(0%);
        }
        .brand-name {
            font-weight: 600;
            color: #000000; /* Pure black in light mode */
        }
        [data-theme="dark"] .brand-name {
            color: #ffffff !important;
        }
        @media (prefers-color-scheme: dark) {
            .brand-name {
                color: #ffffff !important;
            }
        }
        .block-container {
            padding-bottom: 150px !important;
            padding-top: 3rem !important;
        }

        /* --- Theme-Aware Dynamic Styling --- */
        
        /* Universal Typography */
        h1, h2, h3, h4, h5, h6, p, span, label, .stMarkdown {
            color: var(--text-color) !important;
        }

        /* Metric Styling - Fully Dynamic */
        div[data-testid="stMetric"] { 
            background-color: var(--secondary-background-color); 
            padding: 24px; 
            border-radius: 16px; 
            box-shadow: 0 4px 15px rgba(0,0,0,0.05);
            border: 1px solid rgba(128, 128, 128, 0.1);
            transition: all 0.3s ease;
        }
        
        div[data-testid="stMetric"]:hover {
            transform: translateY(-4px);
            box-shadow: 0 8px 25px rgba(0,0,0,0.1);
            border-color: var(--primary-color);
        }

        [data-testid="stMetricLabel"] {
            font-weight: 600 !important;
            opacity: 0.8;
            color: var(--text-color) !important;
        }

        [data-testid="stMetricValue"], 
        [data-testid="stMetricValue"] > div {
            font-weight: 800 !important;
            font-size: 1.6rem !important;
            color: var(--text-color) !important;
            white-space: nowrap !important;
            overflow: visible !important;
            text-overflow: clip !important;
        }

        /* UI Components - Inputs & Selectboxes */
        div[data-testid="stSelectbox"] > div, 
        div[data-testid="stTextInput"] > div,
        div[data-testid="stTextArea"] > div {
            background-color: var(--secondary-background-color) !important;
            color: var(--text-color) !important;
            border-radius: 12px !important;
            border: 1px solid rgba(128, 128, 128, 0.2) !important;
        }

        /* Buttons - High Contrast Primary */
        .stButton>button { 
            width: 100%; 
            border-radius: 12px; 
            height: 3.5em; 
            background-color: var(--primary-color) !important;
            color: white !important; /* Buttons usually have white text regardless of theme */
            font-weight: 700;
            border: none;
            transition: all 0.2s ease;
        }
        .stButton>button:hover {
            opacity: 0.9;
            transform: scale(1.01);
        }

        /* Tabs & Dividers */
        div[data-testid="stTabs"] button {
            color: var(--text-color) !important;
            opacity: 0.7;
        }
        div[data-testid="stTabs"] button[aria-selected="true"] {
            color: var(--primary-color) !important;
            opacity: 1;
            border-bottom: 2px solid var(--primary-color) !important;
        }
        hr { border-color: rgba(128, 128, 128, 0.2) !important; }

        /* Expander & Notifications */
        div[data-testid="stExpander"], div[data-testid="stNotification"] {
            background-color: var(--secondary-background-color) !important;
            border: 1px solid rgba(128, 128, 128, 0.1) !important;
            border-radius: 14px !important;
            color: var(--text-color) !important;
        }

        /* Sticky Footer - Dynamic Contrast & Always Fixed */
        .sticky-footer {
            position: fixed !important;
            left: 0 !important;
            bottom: 0 !important;
            right: 0 !important;
            width: 100% !important;
            background-color: var(--background-color) !important;
            color: var(--text-color) !important;
            text-align: center;
            padding: 20px 0;
            border-top: 1px solid rgba(128, 128, 128, 0.2);
            z-index: 999999;
            font-size: 0.9rem;
            backdrop-filter: blur(10px);
            box-shadow: 0 -5px 25px rgba(0,0,0,0.1);
        }
        
        [data-theme="dark"] .sticky-footer {
            background-color: #0e1117 !important; /* Pure black/dark for contrast */
            box-shadow: 0 -5px 25px rgba(0,0,0,0.4);
        }

        /* Force Whitish text specifically in Dark Mode */
        [data-theme="dark"] .sticky-footer, 
        [data-theme="dark"] .sticky-footer span,
        [data-theme="dark"] .brand-name {
            color: #f8f9fa !important;
        }
        
        @media (prefers-color-scheme: dark) {
            .sticky-footer, .brand-name {
                color: #f8f9fa !important;
            }
        }

        .brand-name {
            font-weight: 700;
            color: var(--primary-color);
        }

        /* --- Responsive Design --- */
        @media (max-width: 768px) {
            .block-container {
                padding-left: 1rem !important;
                padding-right: 1rem !important;
                padding-bottom: 220px !important;
            }
            .footer-content-inner {
                flex-direction: column;
                gap: 8px;
            }
            [data-testid="stMetricValue"] {
                font-size: 1.3rem !important;
            }
            /* Keep footer fixed and prominent even on mobile */
            .sticky-footer {
                padding: 15px 5px !important;
            }
        }
        """, unsafe_allow_html=True)

# --- Helper Functions ---

def log_event(event_type, details):
    """Logs system events to JSON."""
    log_file = os.path.join(FEEDBACK_DIR, "system_logs.json")
    entry = {
        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "type": event_type,
        "details": details
    }
    try:
        logs = []
        if os.path.exists(log_file):
            with open(log_file, "r") as f: logs = json.load(f)
        logs.append(entry)
        with open(log_file, "w") as f: json.dump(logs[-100:], f, indent=4)
    except: pass

def get_product_category(name, mapping_dict=None):
    """Categorizes product based on keywords."""
    if mapping_dict is None:
        mapping_dict = load_category_mapping()
    name_str = str(name).lower()
    for cat, keywords in mapping_dict.items():
        if any(kw.lower() in name_str for kw in keywords):
            return cat
    
    # Special handling for T-Shirts and Shirts
    fs_keywords = ['full sleeve', 'long sleeve', 'fs', 'l/s']
    is_fs = any(kw in name_str for kw in fs_keywords)
    
    if any(kw in name_str for kw in ['t-shirt', 't shirt', 'tee']):
        return 'FS T-Shirt' if is_fs else 'T-Shirt'
    if 'shirt' in name_str:
        return 'FS Shirt' if is_fs else 'HS Shirt'
        
    return 'Others'

def find_columns(df):
    """Auto-detects columns from dataframe."""
    mapping = {
        'name': ['item name', 'product name', 'product', 'item', 'title', 'description', 'name'],
        'cost': ['item cost', 'price', 'unit price', 'cost', 'rate', 'mrp', 'selling price'],
        'qty': ['quantity', 'qty', 'units', 'sold', 'count', 'total quantity'],
        'date': ['date', 'order date', 'month', 'time', 'created at'],
        'order_id': ['order id', 'order #', 'invoice number', 'invoice #', 'order number', 'transaction id', 'id'],
        'phone': ['phone', 'contact', 'mobile', 'cell', 'phone number', 'customer phone']
    }
    found = {}
    actual_cols = list(df.columns)
    lower_cols = [str(c).strip().lower() for c in actual_cols]
    
    for key, aliases in mapping.items():
        # Exact match
        for alias in aliases:
            if alias in lower_cols:
                found[key] = actual_cols[lower_cols.index(alias)]
                break
        # Partial match
        if key not in found:
            for i, col in enumerate(lower_cols):
                if any(alias in col for alias in aliases):
                    found[key] = actual_cols[i]
                    break
    return found

def clean_numeric_series(series):
    """Fast vectorized numeric cleaning removing currency, commas, and whitespace."""
    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series.fillna(0), errors='coerce').fillna(0)
    cleaned = series.astype(str).str.replace(r'[^0-9.]', '', regex=True)
    return pd.to_numeric(cleaned, errors='coerce').fillna(0)

@st.cache_data(show_spinner="Loading sales data...")
def load_data(file_bytes, filename):
    """Loads CSV or Excel data from uploaded bytes with caching."""
    buf = BytesIO(file_bytes)
    if filename.lower().endswith('.csv'):
        return pd.read_csv(buf)
    else:
        return pd.read_excel(buf)

@st.cache_data(show_spinner="Processing analytics & categorizing products...")
def process_analytics(df, mapping):
    """Core data processing and metric calculation."""
    df = df.copy()
    
    # 1. Clean Data (Vectorized)
    df['Clean_Name'] = df[mapping['name']].fillna('Unknown').astype(str)
    exclude_keywords = 'choose any|shoose any|combo|bundle|bundel'
    df = df[~df['Clean_Name'].str.contains(exclude_keywords, case=False, na=False)]

    df['Clean_Cost'] = clean_numeric_series(df[mapping['cost']])
    df['Clean_Qty'] = clean_numeric_series(df[mapping['qty']])
    df.loc[df['Clean_Qty'] < 0, 'Clean_Qty'] = 0
    df['Total Amount'] = df['Clean_Cost'] * df['Clean_Qty']
    
    # Fast unique-name categorization
    active_mapping = load_category_mapping()
    unique_names = df['Clean_Name'].unique()
    cat_lookup = {name: get_product_category(name, active_mapping) for name in unique_names}
    df['Category'] = df['Clean_Name'].map(cat_lookup).fillna('Others')
    
    # 2. Timeframe Detection
    timeframe = ""
    if mapping.get('date') and mapping['date'] in df.columns:
        try:
            dates = pd.to_datetime(df[mapping['date']], errors='coerce').dropna()
            if not dates.empty:
                if dates.dt.to_period('M').nunique() == 1:
                    timeframe = dates.iloc[0].strftime("%B_%Y")
                else:
                    timeframe = f"{dates.min().strftime('%d%b')}_to_{dates.max().strftime('%d%b_%y')}"
        except: timeframe = "Report"

    # 3. Aggregations
    summary = df.groupby('Category').agg({'Clean_Qty': 'sum', 'Total Amount': 'sum'}).reset_index()
    summary.columns = ['Category', 'Total Qty', 'Total Amount']
    
    t_rev = summary['Total Amount'].sum()
    t_qty = summary['Total Qty'].sum()
    
    drilldown = df.groupby(['Category', 'Clean_Cost']).agg({'Clean_Qty': 'sum', 'Total Amount': 'sum'}).reset_index()
    drilldown.columns = ['Category', 'Price', 'Total Qty', 'Total Amount']
    
    top_items = df.groupby('Clean_Name').agg({'Clean_Qty': 'sum', 'Total Amount': 'sum', 'Category': 'first'}).reset_index()
    top_items.columns = ['Product Name', 'Total Qty', 'Total Amount', 'Category']
    top_items = top_items.sort_values('Total Amount', ascending=False)

    product_summary = df.groupby(['Category', 'Clean_Name']).agg({'Clean_Qty': 'sum', 'Total Amount': 'sum'}).reset_index()
    product_summary.columns = ['Category', 'Product Name', 'Total Qty', 'Total Amount']
    
    # 4. Basket Metrics
    avg_basket_value = 0
    group_cols = [c for c in [mapping.get('order_id'), mapping.get('phone')] if c and c in df.columns]
    
    if group_cols:
        order_groups = df.groupby(group_cols).agg({'Total Amount': 'sum'})
        avg_basket_value = order_groups['Total Amount'].mean()
        
    # 5. Unmapped items analysis ("Others")
    unmapped_df = df[df['Category'] == 'Others']
    if not unmapped_df.empty:
        unmapped_items = unmapped_df.groupby('Clean_Name').agg({
            'Clean_Qty': 'sum',
            'Total Amount': 'sum'
        }).reset_index()
        unmapped_items.columns = ['Product Name', 'Total Qty', 'Total Amount']
        unmapped_items = unmapped_items.sort_values('Total Amount', ascending=False)
    else:
        unmapped_items = pd.DataFrame(columns=['Product Name', 'Total Qty', 'Total Amount'])

    unmapped_qty = unmapped_items['Total Qty'].sum() if not unmapped_items.empty else 0
    unmapped_rev = unmapped_items['Total Amount'].sum() if not unmapped_items.empty else 0

    return {
        'drilldown': drilldown,
        'summary': summary,
        'top_items': top_items,
        'product_summary': product_summary,
        'unmapped_items': unmapped_items,
        'unmapped_qty': unmapped_qty,
        'unmapped_rev': unmapped_rev,
        'timeframe': timeframe,
        'avg_basket_value': avg_basket_value,
        'total_qty': t_qty,
        'total_rev': t_rev,
        'total_orders': len(order_groups) if group_cols else 0
    }

# --- UI Components ---

@st.fragment
def render_sidebar(logo_b64=None):
    with st.sidebar:
        if not logo_b64:
            logo_b64 = load_logo()
        logo_html = f'<img src="data:image/png;base64,{logo_b64}" style="width: 40px; height: 40px; border-radius: 8px; object-fit: contain; flex-shrink: 0;">' if logo_b64 else ''
        st.markdown(
            f"""
            <a href="https://deencommerce.com" target="_blank" style="text-decoration: none; color: inherit;">
                <div style="display: flex; align-items: center; gap: 12px; margin-bottom: 8px; padding-top: 2px;">
                    {logo_html}
                    <div style="display: flex; flex-direction: column;">
                        <div style="display: flex; align-items: center; gap: 8px;">
                            <span style="font-size: 1.2rem; font-weight: 700; color: var(--primary-color, #007bff); letter-spacing: -0.01em; line-height: 1.2;">DEEN Commerce</span>
                            <span style="font-size: 0.65rem; background: rgba(0, 123, 255, 0.12); color: var(--primary-color, #007bff); padding: 2px 7px; border-radius: 8px; font-weight: 600; text-transform: uppercase;">Suite</span>
                        </div>
                        <div style="font-size: 0.78rem; color: var(--text-color); opacity: 0.65; font-weight: 400; margin-top: 2px;">
                            Category Intelligence & Performance
                        </div>
                    </div>
                </div>
            </a>
            """,
            unsafe_allow_html=True
        )
        st.divider()

        st.header(":material/feedback: Feedback & Debug")
        comment = st.text_area("Report Issues:", placeholder="Category 'Polo' is incorrect...")
        if st.button("Submit Report", icon=":material/send:"):
            feedback_file = os.path.join(FEEDBACK_DIR, "user_feedback.json")
            entry = {"timestamp": datetime.now().isoformat(), "comment": comment}
            try:
                data = []
                if os.path.exists(feedback_file):
                    with open(feedback_file, "r") as f: data = json.load(f)
                data.append(entry)
                with open(feedback_file, "w") as f: json.dump(data, f, indent=4)
                st.success("Feedback saved!")
            except: st.error("Failed to save.")
        
        st.divider()
        if st.checkbox("View System Logs"):
            log_path = os.path.join(FEEDBACK_DIR, "system_logs.json")
            if os.path.exists(log_path):
                with open(log_path, "r") as f: st.json(json.load(f)[-10:])
            else: st.info("No logs available.")

def render_footer(logo_b64):
    footer_html = f"""
    <div style="height: 120px;"></div> <!-- Spacer to prevent content overlap -->
    <div class="sticky-footer">
        <div class="footer-content-inner">
            <span>© {datetime.now().year} Sajid Islam. All rights reserved. | </span>
            <div class="brand-wrapper">
                <span style="opacity: 0.85;">Powered by</span>
                <img src="data:image/png;base64,{logo_b64}" class="small-logo">
                <span class="brand-name">DEEN Commerce Ltd.</span>
            </div>
        </div>
    </div>
    """
    st.markdown(footer_html, unsafe_allow_html=True)

def generate_excel_report(sheet_dict):
    """Generates an optimized, beautifully styled Excel workbook with fast row iteration."""
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        for sheet_name, dframe in sheet_dict.items():
            dframe.to_excel(writer, sheet_name=sheet_name, index=False)
        
        wb = writer.book
        header_fill = PatternFill(start_color='007BFF', end_color='007BFF', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        total_fill = PatternFill(start_color='ECECEC', end_color='ECECEC', fill_type='solid')
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        
        for sheet_name, dframe in sheet_dict.items():
            ws = wb[sheet_name]
            ws.freeze_panes = 'A2'
            
            # Style Header Row
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = thin_border
            
            # Pre-compute column configurations
            col_alignments = [
                Alignment(horizontal='left') if col in ['Category', 'Product Name', 'Product'] 
                else Alignment(horizontal='right')
                for col in dframe.columns
            ]
            is_numeric_col = [any(x in col for x in ['Qty', 'Price', 'Amount']) for col in dframe.columns]
            
            # Auto-adjust column widths
            for i, col_name in enumerate(dframe.columns):
                col_letter = ws.cell(row=1, column=i+1).column_letter
                content_len = dframe[col_name].astype(str).str.len().max() if not dframe.empty else 0
                ws.column_dimensions[col_letter].width = max(int(content_len), len(str(col_name))) + 3
            
            # Fast single-pass row styling using iter_rows
            for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(dframe.columns)):
                first_cell_str = str(row_cells[0].value).upper()
                is_total = 'TOTAL' in first_cell_str
                for idx, cell in enumerate(row_cells):
                    cell.border = thin_border
                    cell.alignment = col_alignments[idx]
                    if is_numeric_col[idx]:
                        cell.number_format = '#,##0'
                    if is_total:
                        cell.font = Font(bold=True)
                        cell.fill = total_fill
    buf.seek(0)
    return buf.getvalue()

# --- Main App ---

def main():
    st.set_page_config(page_title="Sales Dashboard", page_icon=":material/monitoring:", layout="wide")
    apply_custom_styles()
    
    logo_b64 = load_logo()
    render_sidebar(logo_b64)
    
    st.title(":material/monitoring: Sales Performance Dashboard")
    
    uploaded_file = st.file_uploader("Upload Sales Data (Excel or CSV)", type=['xlsx', 'csv'])
    
    if uploaded_file:
        try:
            # Check if file changed to reset state
            file_key = f"{uploaded_file.name}_{uploaded_file.size}"
            if st.session_state.get('current_file_key') != file_key:
                st.session_state['current_file_key'] = file_key
                st.session_state['results'] = None

            file_bytes = uploaded_file.getvalue()
            df = load_data(file_bytes, uploaded_file.name)
            st.success(f"Attached: {uploaded_file.name}")
            
            with st.expander("Preview Data", icon=":material/visibility:", expanded=False):
                st.dataframe(df.head(10), width="stretch")

            # Column Mapping Section
            auto_cols = find_columns(df)
            all_cols = list(df.columns)
            mandatory_keys = ['name', 'cost', 'qty']
            is_mapped = all(k in auto_cols for k in mandatory_keys)
            
            def get_idx(key):
                return all_cols.index(auto_cols[key]) if key in auto_cols else 0

            with st.container(border=True):
                st.subheader(":material/tune: Verify Column Mapping", anchor=False)
                if is_mapped:
                    st.caption(":material/check_circle: Mandatory columns were auto-detected. Adjust if needed:")
                else:
                    st.caption(":material/info: Please match the required columns from your uploaded file:")

                st.markdown("**:material/check_circle: Required Fields** *(Mandatory for calculations)*")
                rc1, rc2, rc3 = st.columns(3)
                m_name = rc1.selectbox("Product Name *", all_cols, index=get_idx('name'), help="Column containing the product or item description")
                m_cost = rc2.selectbox("Unit Price *", all_cols, index=get_idx('cost'), help="Column containing unit price or cost")
                m_qty = rc3.selectbox("Quantity *", all_cols, index=get_idx('qty'), help="Column containing quantity sold")

                st.markdown("**:material/add_circle: Optional Dimensions** *(Enables timeframe & basket size metrics)*")
                oc1, oc2, oc3 = st.columns(3)
                m_date = oc1.selectbox("Date (Opt)", ["None"] + all_cols, index=get_idx('date')+1 if 'date' in auto_cols else 0, help="Order date column for report timeframe")
                m_order = oc2.selectbox("Order ID (Opt)", ["None"] + all_cols, index=get_idx('order_id')+1 if 'order_id' in auto_cols else 0, help="Order number or invoice ID for basket size")
                m_phone = oc3.selectbox("Phone (Opt)", ["None"] + all_cols, index=get_idx('phone')+1 if 'phone' in auto_cols else 0, help="Customer contact for multi-item order grouping")

                mapping = {
                    'name': m_name, 'cost': m_cost, 'qty': m_qty,
                    'date': m_date if m_date != "None" else None,
                    'order_id': m_order if m_order != "None" else None,
                    'phone': m_phone if m_phone != "None" else None
                }

                st.write("")
                if st.button("Generate Analytics", icon=":material/analytics:", type="primary"):
                    st.session_state['results'] = process_analytics(df, mapping)

            results = st.session_state.get('results')
            if results is not None:
                # Metrics Row
                m1, m2, m3, m4 = st.columns(4)
                m1.metric("Total Orders", f"{results['total_orders']:,.0f}", border=True)
                m2.metric("Units Sold", f"{results['total_qty']:,.0f}", border=True)
                m3.metric("Gross Revenue", f"TK {results['total_rev']:,.0f}", border=True)
                
                # Show Basket Value if data available
                if results['avg_basket_value'] > 0:
                    m4.metric("Basket Size", f"TK {results['avg_basket_value']:,.0f}", border=True, help="Average gross spend per order or customer phone")
                else:
                    m4.metric("Basket Size", "N/A", border=True, help="Order ID or Phone column not provided")

                st.divider()
                
                # Visuals
                v1, v2 = st.columns(2)
                
                # Sort revenue for pie chart and quantity for bar chart (big to small)
                summ_sorted = results['summary'].sort_values('Total Amount', ascending=False)
                summ_qty_sorted = results['summary'].sort_values('Total Qty', ascending=False)
                
                # Synchronized color map so each category has the exact same color in both charts
                cat_color_map = get_category_color_map(results['summary']['Category'].unique())
                
                with v1:
                    with st.container(border=True):
                        st.subheader(":material/pie_chart: Revenue by Category", anchor=False)
                        fig_pie = px.pie(
                            summ_sorted, 
                            values='Total Amount', 
                            names='Category', 
                            color='Category', 
                            hole=0.5, 
                            color_discrete_map=cat_color_map
                        )
                        fig_pie.update_layout(
                            paper_bgcolor='rgba(0,0,0,0)',
                            plot_bgcolor='rgba(0,0,0,0)',
                            margin=dict(t=10, b=10, l=10, r=10),
                            legend=dict(orientation="v", yanchor="middle", y=0.5, xanchor="left", x=1.0)
                        )
                        st.plotly_chart(fig_pie, width="stretch")
                
                with v2:
                    with st.container(border=True):
                        st.subheader(":material/bar_chart: Volume by Category", anchor=False)
                        fig_bar = px.bar(
                            summ_qty_sorted, 
                            x='Category', 
                            y='Total Qty', 
                            color='Category', 
                            color_discrete_map=cat_color_map
                        )
                        fig_bar.update_layout(
                            xaxis={'categoryorder': 'total descending'},
                            paper_bgcolor='rgba(0,0,0,0)',
                            plot_bgcolor='rgba(0,0,0,0)',
                            margin=dict(t=10, b=10, l=10, r=10),
                            showlegend=False
                        )
                        st.plotly_chart(fig_bar, width="stretch")
                
                # Data Tables
                t1, t2, t3 = st.tabs([":material/table_chart: Category Summary", ":material/payments: Price-wise Category", ":material/inventory_2: Category-wise Product"])
                
                df_breakdown = results['summary'].sort_values('Category', ascending=True).copy()
                
                # Add Total Sales Row
                total_qty = df_breakdown['Total Qty'].sum()
                total_amount = df_breakdown['Total Amount'].sum()
                
                # Append row directly (using loc to avoid append/concat warnings)
                df_breakdown.loc[len(df_breakdown) + 1] = ['TOTAL SALES (Summary)', total_qty, total_amount]
                
                df_breakdown.index = range(1, len(df_breakdown) + 1)
                
                df_drill = results['drilldown'].sort_values(['Category', 'Price'], ascending=[True, False]).copy()
                df_drill.columns = ['Category', 'Price', 'Qty', 'Total Amount']
                
                # Add Total Row
                df_drill.loc[len(df_drill) + 1] = ['TOTAL (All Categories)', None, df_drill['Qty'].sum(), df_drill['Total Amount'].sum()]
                df_drill.index = range(1, len(df_drill) + 1)

                df_product = results['product_summary'].sort_values(['Category', 'Total Amount'], ascending=[True, False]).copy()
                df_product.columns = ['Category', 'Product Name', 'Total Qty', 'Total Amount']
                
                # Add Total Row
                df_product.loc[len(df_product) + 1] = ['TOTAL (All Products)', None, df_product['Total Qty'].sum(), df_product['Total Amount'].sum()]
                df_product.index = range(1, len(df_product) + 1)

                with t1: 
                    st.dataframe(
                        df_breakdown[['Category', 'Total Qty', 'Total Amount']],
                        column_config={
                            "Category": st.column_config.TextColumn("Product Category"),
                            "Total Qty": st.column_config.NumberColumn("Total Units", format="%,.0f pcs"),
                            "Total Amount": st.column_config.NumberColumn("Total Revenue", format="TK %,.0f"),
                        },
                        width="stretch",
                        hide_index=True
                    )
                
                with t2: 
                    st.dataframe(
                        df_drill[['Category', 'Price', 'Qty', 'Total Amount']],
                        column_config={
                            "Category": st.column_config.TextColumn("Product Category"),
                            "Price": st.column_config.NumberColumn("Unit Price", format="TK %,.0f"),
                            "Qty": st.column_config.NumberColumn("Units Sold", format="%,.0f pcs"),
                            "Total Amount": st.column_config.NumberColumn("Total Revenue", format="TK %,.0f"),
                        },
                        width="stretch",
                        hide_index=True
                    )

                with t3:
                    c_search, c_filter = st.columns([3, 2])
                    with c_search:
                        p_search = st.text_input(
                            "Search Product", 
                            placeholder="Search products by title or keyword...", 
                            label_visibility="collapsed",
                            key="tab3_product_search"
                        )
                    with c_filter:
                        cat_choices = sorted([c for c in results['product_summary']['Category'].unique() if c])
                        selected_cats = st.multiselect(
                            "Filter Category", 
                            options=cat_choices, 
                            default=[], 
                            placeholder="All Categories",
                            label_visibility="collapsed",
                            key="tab3_category_filter"
                        )

                    filtered_prod = results['product_summary'].copy()
                    if selected_cats:
                        filtered_prod = filtered_prod[filtered_prod['Category'].isin(selected_cats)]
                    if p_search.strip():
                        filtered_prod = filtered_prod[
                            filtered_prod['Product Name'].str.contains(p_search.strip(), case=False, na=False)
                        ]

                    df_filtered_view = filtered_prod.sort_values(['Category', 'Total Amount'], ascending=[True, False]).copy()
                    df_filtered_view.columns = ['Category', 'Product Name', 'Total Qty', 'Total Amount']
                    
                    f_total_qty = df_filtered_view['Total Qty'].sum()
                    f_total_amt = df_filtered_view['Total Amount'].sum()
                    
                    df_filtered_view.loc[len(df_filtered_view) + 1] = [
                        f"TOTAL ({len(df_filtered_view)} items)", None, f_total_qty, f_total_amt
                    ]
                    df_filtered_view.index = range(1, len(df_filtered_view) + 1)

                    st.dataframe(
                        df_filtered_view[['Category', 'Product Name', 'Total Qty', 'Total Amount']],
                        column_config={
                            "Category": st.column_config.TextColumn("Category"),
                            "Product Name": st.column_config.TextColumn("Product Name"),
                            "Total Qty": st.column_config.NumberColumn("Total Units", format="%,.0f pcs"),
                            "Total Amount": st.column_config.NumberColumn("Total Revenue", format="TK %,.0f"),
                        },
                        width="stretch",
                        hide_index=True
                    )

                # Unmapped Items Inspector ("Others" Review)
                unmapped_df = results.get('unmapped_items')
                unmapped_count = len(unmapped_df) if unmapped_df is not None else 0
                with st.expander(f"Unmapped Products Inspector ('Others' — {unmapped_count} unique items)", icon=":material/warning:", expanded=False):
                    if unmapped_df is None or unmapped_df.empty:
                        st.success("All products were successfully categorized! Zero unmapped items.", icon=":material/check_circle:")
                    else:
                        u1, u2 = st.columns(2)
                        u1.metric("Unmapped Units", f"{results['unmapped_qty']:,.0f}", border=True)
                        u2.metric("Unmapped Revenue", f"TK {results['unmapped_rev']:,.0f}", border=True)
                        st.caption("Review these products to identify missing keywords for your category rules in `config/categories.json`:")
                        st.dataframe(
                            unmapped_df,
                            column_config={
                                "Product Name": st.column_config.TextColumn("Product Description"),
                                "Total Qty": st.column_config.NumberColumn("Quantity", format="%,.0f pcs"),
                                "Total Amount": st.column_config.NumberColumn("Revenue Impact", format="TK %,.0f"),
                            },
                            width="stretch",
                            hide_index=True
                        )

                # Export
                sheet_dict = {
                    'Category Summary': df_breakdown,
                    'Price-wise Category': df_drill,
                    'Category-wise Product': df_product
                }

                excel_data = generate_excel_report(sheet_dict)
                fname = f"Sales_Report_{results['timeframe']}.xlsx"
                st.download_button("Download Report", data=excel_data, file_name=fname, icon=":material/download:")
                
        except Exception as e:
            st.error(f"Processing Error: {e}")
            log_event("CRASH", str(e))
    else:
        st.session_state['current_file_key'] = None
        st.session_state['results'] = None

    render_footer(logo_b64)

if __name__ == "__main__":
    main()
